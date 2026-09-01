"""Excel 模板编辑对话框 — 3 模式 + 单行布局

布局:
- 模式 1:配置 sheet / 英文表名列 / Sheet 名称列  三个 QComboBox 在一行
- 模式 2:不显示
- 模式 3:sheet 列表 + 英文表名(从 project db_table.name 选,纯下拉)
- 数据页:表头行 / 数据起始行 / 数据起始列  三个 QSpinBox 在一行
"""
from __future__ import annotations
from typing import Optional, Iterable
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QTextEdit,
    QPushButton, QSpinBox, QComboBox, QFormLayout, QFrame, QTableWidget,
    QTableWidgetItem, QFileDialog, QHeaderView, QAbstractItemView,
    QMessageBox,
)
import qtawesome as qta
import json
import csv
import openpyxl

from app.ui.i18n import tr
from app.repos.excel_template_repo import ExcelTemplate
from app.repos.project_repo import Project


PARSE_MODES = [
    ("mapping",      "模式 1:配置表(TOTAL + 英文表名列 + Sheet 名称列)"),
    ("sheet_name",   "模式 2:sheet 名 = 英文表名"),
    ("chinese_name", "模式 3:sheet 是中文,用映射转英文表名"),
]


def _col_letter(idx: int) -> str:
    """0-based index → Excel 字母 (0=A, 25=Z, 26=AA)"""
    n = idx + 1
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(ord('A') + r) + s
    return s


class ExcelTemplateDialog(QDialog):
    def __init__(self, template: Optional[ExcelTemplate] = None,
                 projects: Optional[Iterable[Project]] = None,
                 file_path: str = "",
                 project_id: Optional[int] = None,
                 project_tables: Optional[list[str]] = None,
                 parent=None):
        super().__init__(parent)
        self._template = template
        self._file_path = file_path
        self._project_id = project_id
        self._project_tables = project_tables or []
        self._sheet_names: list[str] = []
        self._config_cols_count: int = 0  # config sheet 的列数(决定字母下拉范围)
        self.setWindowTitle(tr("dlg.excel.new") if template is None else tr("dlg.excel.edit"))
        self.setMinimumWidth(720)
        self.resize(820, 720)
        self._build(projects)
        if self._file_path:
            self._load_sheet_names()
            self._restore_template_state()

    def _build(self, projects):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(10)

        # =============== 基础信息 ===============
        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        form.setHorizontalSpacing(12)
        form.setVerticalSpacing(8)

        self.name_edit = QLineEdit()
        if self._template:
            self.name_edit.setText(self._template.template_name)
        form.addRow(tr("dlg.excel.name") + " *", self.name_edit)

        self.project_combo = QComboBox()
        self.project_combo.addItem(tr("sqllib.project_filter_global"), None)
        if projects:
            for p in projects:
                self.project_combo.addItem(p.name, p.id)
        if self._template and self._template.project_id is not None:
            idx = self.project_combo.findData(self._template.project_id)
            if idx >= 0:
                self.project_combo.setCurrentIndex(idx)
        elif self._project_id is not None:
            idx = self.project_combo.findData(self._project_id)
            if idx >= 0:
                self.project_combo.setCurrentIndex(idx)
        self.project_combo.currentIndexChanged.connect(self._on_project_changed)
        form.addRow(tr("dlg.excel.project"), self.project_combo)

        self.mode_combo = QComboBox()
        for code, label in PARSE_MODES:
            self.mode_combo.addItem(label, code)
        if self._template:
            for i, (code, _) in enumerate(PARSE_MODES):
                if code == self._template.parse_mode:
                    self.mode_combo.setCurrentIndex(i)
                    break
        self.mode_combo.currentIndexChanged.connect(self._on_mode_changed)
        form.addRow("解析模式 *", self.mode_combo)

        # Excel 文件
        file_row = QHBoxLayout()
        self.path_edit = QLineEdit()
        self.path_edit.setText(self._file_path)
        self.path_edit.setPlaceholderText("选 xlsx → 加载 sheet 列表")
        file_row.addWidget(self.path_edit, 1)
        browse_btn = QPushButton("浏览")
        browse_btn.setObjectName("Ghost")
        browse_btn.clicked.connect(self._on_browse_file)
        file_row.addWidget(browse_btn)
        reload_btn = QPushButton("刷新 sheet")
        reload_btn.setObjectName("Ghost")
        reload_btn.clicked.connect(self._reload_sheet_names)
        file_row.addWidget(reload_btn)
        form.addRow("Excel 文件", self._wrap(file_row))

        layout.addLayout(form)

        # =============== 模式 1:三个字段一行 ===============
        self.mode1_frame = QFrame()
        self.mode1_frame.setObjectName("Card")
        m1v = QVBoxLayout(self.mode1_frame)
        m1v.setContentsMargins(12, 12, 12, 12)
        m1v.setSpacing(8)

        # 标题 + hint
        m1h = QHBoxLayout()
        m1h.addWidget(QLabel("配置(TOTAL 页)"))
        m1h.addStretch()
        m1h.addWidget(QLabel("💡 先选 xlsx + config sheet,字母下拉自动出"))
        m1v.addLayout(m1h)

        # 三个下拉在一行
        self.sheet_combo = QComboBox()
        self.sheet_combo.setPlaceholderText("—")
        self.sheet_combo.currentTextChanged.connect(self._on_config_sheet_changed)
        self.tn_col_combo = QComboBox()
        self.tn_col_combo.setPlaceholderText("—")
        self.sn_col_combo = QComboBox()
        self.sn_col_combo.setPlaceholderText("—")

        row1 = QHBoxLayout()
        row1.setSpacing(8)
        row1.addWidget(QLabel("配置 sheet *"))
        row1.addWidget(self.sheet_combo, 1)
        row1.addWidget(QLabel("英文表名列 *"))
        row1.addWidget(self.tn_col_combo, 1)
        row1.addWidget(QLabel("Sheet 名称列 *"))
        row1.addWidget(self.sn_col_combo, 1)
        m1v.addLayout(row1)

        layout.addWidget(self.mode1_frame)

        # =============== 模式 3:映射表 ===============
        self.mode3_frame = QFrame()
        self.mode3_frame.setObjectName("Card")
        m3v = QVBoxLayout(self.mode3_frame)
        m3v.setContentsMargins(12, 12, 12, 12)
        m3v.setSpacing(8)

        m3h = QHBoxLayout()
        m3h.addWidget(QLabel("Sheet 列表 → 英文表名(从项目表结构里选)"))
        m3h.addStretch()
        import_csv_btn = QPushButton("导入 CSV")
        import_csv_btn.setObjectName("Ghost")
        import_csv_btn.clicked.connect(self._on_import_csv)
        m3h.addWidget(import_csv_btn)
        import_json_btn = QPushButton("导入 JSON")
        import_json_btn.setObjectName("Ghost")
        import_json_btn.clicked.connect(self._on_import_json)
        m3h.addWidget(import_json_btn)
        m3v.addLayout(m3h)

        self.mapping_table = QTableWidget(0, 2)
        self.mapping_table.setHorizontalHeaderLabels(["Sheet 名", "英文表名(下拉选)"])
        self.mapping_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.mapping_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        # 关键:加大行高,让 cell widget 的 combo 能完整显示(下拉文字 + ▼ 箭头)
        self.mapping_table.verticalHeader().setDefaultSectionSize(40)
        # 整张表给个最低高度
        self.mapping_table.setMinimumHeight(280)
        m3v.addWidget(self.mapping_table, 1)

        hint3 = QLabel("💡 不填表名的 sheet 不解析")
        hint3.setObjectName("Muted")
        m3v.addWidget(hint3)

        layout.addWidget(self.mode3_frame)

        # =============== 数据页配置:三个字段一行 ===============
        data_frame = QFrame()
        data_frame.setObjectName("Card")
        dfv = QVBoxLayout(data_frame)
        dfv.setContentsMargins(12, 8, 12, 8)
        dfv.setSpacing(6)
        dfv.addWidget(QLabel("数据页配置(所有 sheet 共用)"))
        self.header_spin = QSpinBox()
        self.header_spin.setRange(1, 9999)
        self.header_spin.setValue(self._template.header_row if self._template else 1)
        self.data_spin = QSpinBox()
        self.data_spin.setRange(1, 9999)
        self.data_spin.setValue(self._template.data_start_row if self._template else 2)
        self.col_spin = QSpinBox()
        self.col_spin.setRange(1, 9999)
        self.col_spin.setValue(self._template.column_start if self._template else 1)
        row2 = QHBoxLayout()
        row2.setSpacing(8)
        row2.addWidget(QLabel("表头行"))
        row2.addWidget(self.header_spin)
        row2.addWidget(QLabel("数据起始行"))
        row2.addWidget(self.data_spin)
        row2.addWidget(QLabel("数据起始列"))
        row2.addWidget(self.col_spin)
        row2.addStretch()
        dfv.addLayout(row2)
        layout.addWidget(data_frame)

        # =============== 描述 ===============
        layout.addWidget(QLabel(tr("dlg.excel.desc")))
        self.desc_edit = QTextEdit()
        self.desc_edit.setMaximumHeight(50)
        if self._template and self._template.description:
            self.desc_edit.setPlainText(self._template.description)
        layout.addWidget(self.desc_edit)

        # Error
        self.error_label = QLabel("")
        self.error_label.setObjectName("Danger")
        self.error_label.setVisible(False)
        layout.addWidget(self.error_label)

        # Buttons
        btn_row = QHBoxLayout()
        btn_row.addStretch()
        cancel = QPushButton(tr("action.cancel"))
        cancel.setObjectName("Ghost")
        cancel.clicked.connect(self.reject)
        btn_row.addWidget(cancel)
        ok = QPushButton(tr("action.save"))
        ok.setObjectName("Primary")
        ok.clicked.connect(self._on_accept)
        btn_row.addWidget(ok)
        layout.addLayout(btn_row)

        self._on_mode_changed()
        self.name_edit.setFocus()

    def _wrap(self, layout) -> QWidget:
        w = QFrame()
        w.setLayout(layout)
        layout.setContentsMargins(0, 0, 0, 0)
        return w

    def _restore_template_state(self):
        """从已有 template 恢复 UI 状态(老 template 没 sheet 名也能恢复)"""
        if not self._template:
            return
        if self._template.parse_mode == "mapping" and self._template.config_sheet_name:
            if self._template.config_sheet_name in self._sheet_names:
                self.sheet_combo.setCurrentText(self._template.config_sheet_name)
                # 触发 _on_config_sheet_changed
            else:
                # config sheet 在 xlsx 里没找到,加进去
                self.sheet_combo.addItem(self._template.config_sheet_name)
                self.sheet_combo.setCurrentText(self._template.config_sheet_name)
        if self._template.parse_mode == "chinese_name":
            self._load_mapping_from_template(self._template.get_name_mapping())

    # ----- 文件选择 / sheet 加载 -----
    def _on_browse_file(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "选 Excel 文件", self._file_path,
            "Excel files (*.xlsx *.xls);;All files (*)"
        )
        if path:
            self._file_path = path
            self.path_edit.setText(path)
            self._load_sheet_names()

    def _reload_sheet_names(self) -> None:
        if not self._file_path:
            QMessageBox.information(self, "提示", "先选 Excel 文件")
            return
        self._load_sheet_names()

    def _load_sheet_names(self) -> None:
        try:
            wb = openpyxl.load_workbook(self._file_path, read_only=True, data_only=True)
            self._sheet_names = list(wb.sheetnames)
            wb.close()
        except Exception as e:
            QMessageBox.warning(self, "错误", f"无法读取 Excel: {e}")
            self._sheet_names = []
        # 刷 sheet_combo(纯下拉,非 editable)
        self.sheet_combo.blockSignals(True)
        current = self.sheet_combo.currentText()
        self.sheet_combo.clear()
        self.sheet_combo.addItem("")  # 占位
        self.sheet_combo.addItems(self._sheet_names)
        if current and current in self._sheet_names:
            self.sheet_combo.setCurrentText(current)
        self.sheet_combo.blockSignals(False)
        # 刷映射表(模式 3)
        if self.mode_combo.currentData() == "chinese_name":
            self._refresh_mapping_table()

    def _on_config_sheet_changed(self, sheet_name: str) -> None:
        """用户选 config sheet → 读列数,刷两个字母下拉"""
        if not sheet_name or not self._file_path:
            return
        try:
            wb = openpyxl.load_workbook(self._file_path, data_only=True, read_only=True)
            if sheet_name not in wb.sheetnames:
                wb.close()
                return
            ws = wb[sheet_name]
            # 读首行拿列数
            rows = list(ws.iter_rows(values_only=True, max_row=1))
            wb.close()
            if not rows or not rows[0]:
                self._config_cols_count = 0
                return
            self._config_cols_count = len(rows[0])
        except Exception as e:
            QMessageBox.warning(self, "错误", f"读取 config sheet 失败: {e}")
            self._config_cols_count = 0
            return
        # 刷两个字母下拉 (A, B, C, ..., Z, AA, AB, ...)
        letters = [_col_letter(i) for i in range(self._config_cols_count)]
        current_tn = self.tn_col_combo.currentText()
        current_sn = self.sn_col_combo.currentText()
        self.tn_col_combo.clear()
        self.tn_col_combo.addItems(letters)
        if current_tn in letters:
            self.tn_col_combo.setCurrentText(current_tn)
        self.sn_col_combo.clear()
        self.sn_col_combo.addItems(letters)
        if current_sn in letters:
            self.sn_col_combo.setCurrentText(current_sn)

    def _on_project_changed(self, _idx: int) -> None:
        from app.services.registry import reg
        pid = self.project_combo.currentData()
        if pid:
            try:
                tables = reg().table_repo.list_by_project(pid)
                self._project_tables = [t.name for t in tables]
            except Exception:
                self._project_tables = []
        else:
            self._project_tables = []
        if self.mode_combo.currentData() == "chinese_name":
            self._refresh_mapping_table()

    def _on_mode_changed(self):
        mode = self.mode_combo.currentData()
        self.mode1_frame.setVisible(mode == "mapping")
        self.mode3_frame.setVisible(mode == "chinese_name")
        if mode == "chinese_name":
            if self.mapping_table.rowCount() == 0 and self._sheet_names:
                self._refresh_mapping_table()

    # ----- 模式 3:映射表 -----
    def _refresh_mapping_table(self):
        existing = self._collect_mapping()
        self.mapping_table.setRowCount(0)
        for sn in self._sheet_names:
            r = self.mapping_table.rowCount()
            self.mapping_table.insertRow(r)
            sn_item = QTableWidgetItem(sn)
            sn_item.setFlags(sn_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            self.mapping_table.setItem(r, 0, sn_item)
            # 英文表名 — 纯下拉(不可手输),第一项是"不填(不解析)"
            combo = QComboBox()
            combo.setMinimumHeight(32)
            combo.addItem("— 不填(不解析) —", "")
            if self._project_tables:
                combo.addItems(self._project_tables)
            else:
                combo.addItem("(项目无表结构)", "_NONE_")
            if sn in existing:
                idx = combo.findText(existing[sn])
                if idx >= 0:
                    combo.setCurrentIndex(idx)
            self.mapping_table.setCellWidget(r, 1, combo)

    def _load_mapping_from_template(self, mapping: dict):
        for r in range(self.mapping_table.rowCount()):
            sn_item = self.mapping_table.item(r, 0)
            combo = self.mapping_table.cellWidget(r, 1)
            if not sn_item or not isinstance(combo, QComboBox):
                continue
            sn = sn_item.text()
            if sn in mapping:
                idx = combo.findText(mapping[sn])
                if idx >= 0:
                    combo.setCurrentIndex(idx)

    def _collect_mapping(self) -> dict:
        m = {}
        for r in range(self.mapping_table.rowCount()):
            sn_item = self.mapping_table.item(r, 0)
            combo = self.mapping_table.cellWidget(r, 1)
            if not sn_item or not isinstance(combo, QComboBox):
                continue
            sn = sn_item.text().strip()
            tn = combo.currentText().strip()
            # 跳过"不填"和"项目无表结构"占位
            if not tn or tn.startswith("—") or tn.startswith("(") or combo.currentData() in ("", "_NONE_"):
                continue
            if sn and tn:
                m[sn] = tn
        return m

    def _on_import_csv(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "导入 CSV(两列:中文,英文)", "", "CSV files (*.csv);;All files (*)"
        )
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)
            m = {}
            for i, row in enumerate(rows):
                if i == 0 and len(row) >= 2 and row[0] in ("中文", "sheet", "sheet_name", "中文名"):
                    continue
                if len(row) >= 2 and row[0].strip() and row[1].strip():
                    m[row[0].strip()] = row[1].strip()
            matched = 0
            for r in range(self.mapping_table.rowCount()):
                sn_item = self.mapping_table.item(r, 0)
                combo = self.mapping_table.cellWidget(r, 1)
                if sn_item and isinstance(combo, QComboBox):
                    sn = sn_item.text().strip()
                    if sn in m:
                        idx = combo.findText(m[sn])
                        if idx >= 0:
                            combo.setCurrentIndex(idx)
                            matched += 1
            unmatched = [k for k in m if k not in self._sheet_names]
            if unmatched:
                QMessageBox.information(
                    self, "导入完成",
                    f"已合并 {matched} 条映射\n"
                    f"以下 {len(unmatched)} 条不在当前 sheet 列表:\n"
                    + "\n".join(unmatched[:20])
                )
            else:
                QMessageBox.information(self, "导入完成", f"已合并 {matched} 条映射")
        except Exception as e:
            QMessageBox.warning(self, "错误", f"导入失败: {e}")

    def _on_import_json(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "导入 JSON", "", "JSON files (*.json);;All files (*)"
        )
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                m = json.load(f)
            if not isinstance(m, dict):
                raise ValueError("JSON 必须是对象(中文→英文)")
            matched = 0
            for r in range(self.mapping_table.rowCount()):
                sn_item = self.mapping_table.item(r, 0)
                combo = self.mapping_table.cellWidget(r, 1)
                if sn_item and isinstance(combo, QComboBox):
                    sn = sn_item.text().strip()
                    if sn in m:
                        idx = combo.findText(str(m[sn]))
                        if idx >= 0:
                            combo.setCurrentIndex(idx)
                            matched += 1
            QMessageBox.information(self, "导入完成", f"已合并 {matched} 条映射")
        except Exception as e:
            QMessageBox.warning(self, "错误", f"导入失败: {e}")

    def _on_accept(self) -> None:
        if not self.name_edit.text().strip():
            self._show_err("模板名必填")
            return
        mode = self.mode_combo.currentData()
        if mode == "mapping":
            if not self.sheet_combo.currentText().strip():
                self._show_err("模式 1 需要选配置 sheet")
                return
            if not self.tn_col_combo.currentText().strip():
                self._show_err("模式 1 需要选英文表名列")
                return
            if not self.sn_col_combo.currentText().strip():
                self._show_err("模式 1 需要选 sheet 名称列")
                return
        elif mode == "chinese_name":
            m = self._collect_mapping()
            if not m:
                self._show_err("模式 3 至少要填一条 sheet→英文表名映射")
                return
        if self.data_spin.value() < self.header_spin.value():
            self._show_err("数据起始行不能小于表头行")
            return
        self.accept()

    def _show_err(self, msg: str) -> None:
        self.error_label.setText(msg)
        self.error_label.setVisible(True)

    def get_template(self) -> dict:
        mode = self.mode_combo.currentData()
        project_id = self.project_combo.currentData()
        if mode == "chinese_name":
            mapping = json.dumps(self._collect_mapping(), ensure_ascii=False)
        else:
            mapping = "{}"
        return {
            "template_name": self.name_edit.text().strip(),
            "project_id": project_id,
            "parse_mode": mode,
            "config_sheet_name": self.sheet_combo.currentText().strip() if mode == "mapping" else "",
            "table_name_col": self.tn_col_combo.currentText().strip() if mode == "mapping" else "",
            "sheet_name_col": self.sn_col_combo.currentText().strip() if mode == "mapping" else "",
            "header_row": self.header_spin.value(),
            "data_start_row": self.data_spin.value(),
            "column_start": self.col_spin.value(),
            "name_mapping": mapping,
            "description": self.desc_edit.toPlainText().strip(),
        }
